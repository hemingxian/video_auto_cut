# -*- coding: utf-8 -*-
"""批量干跑 + 导入闸门静态预检，结果合并写进 干跑结果.json（不调模型/API，不联网）。

用法：
    VAC_TOOLS=<仓库>/VideoAutoCut/tools python3 -B dryrun_all.py 干跑结果.json ../*.xlsx

- 每个 xlsx：dryrun.run(path, shots)，shots 直接读该文件「分镜脚本」表（表头第 1 行）。
  结果写在键 = 文件名 下，与 2026-09-23 首版 干跑结果.json 的生成口径一致（已核对逐字节相同）。
- 「闸门预检_方案§3.5」：按开发方案 §3.5 的颜色/状态规则静态读 xlsx，列出将来模板导入器会给出的
  阻断（红格、正文四列黄格、时长为空、校验失败、没有分镜）与提醒（制作列黄格、待确认、状态说明行）。
  这不是仓库代码的输出，仓库 1bc0fdb 的 script_card_fields 不读底色，也不读「来源/核对」。
- out.json 已存在时只更新本次涉及的键，其它键（对照_*、逐字出处校验_*）原样保留。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dryrun  # noqa: E402  (VAC_TOOLS 在 dryrun 里读取)

TEXT = ("分镜标题", "口播", "字幕", "上屏文字")
PROD = ("段落", "时长(秒)", "景别", "运镜", "拍摄角度", "画面内容", "卡片提示", "推荐素材来源",
        "素材要求", "搜索关键词", "转场", "备注")
YELLOW, RED = "FFF2CC", "FFC7CE"
GATE_KEY = "闸门预检_方案§3.5"


def story_rows(path):
    ws = openpyxl.load_workbook(path)["分镜脚本"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if all(c.value in (None, "") for c in row):
            continue
        yield {head[c.column - 1]: c for c in row if head[c.column - 1]}


def shots_of(path):
    return [{k: c.value for k, c in r.items()} for r in story_rows(path)]


def gate(path):
    blocks, warns, n_shots = [], [], 0
    for cells in story_rows(path):
        n_shots += 1
        n = cells["镜号"].value
        for name, c in cells.items():
            rgb = (c.fill.fgColor.rgb or "")[-6:] if c.fill is not None and c.fill.fill_type else ""
            if rgb == RED:
                blocks.append({"shot": n, "cell": c.coordinate, "field": name, "code": "red_cell"})
            elif rgb == YELLOW and name in TEXT:
                blocks.append({"shot": n, "cell": c.coordinate, "field": name, "code": "text_cell_pending"})
            elif rgb == YELLOW and name in PROD:
                warns.append({"shot": n, "cell": c.coordinate, "field": name, "code": "ai_suggested"})
        if cells["时长(秒)"].value in (None, ""):
            blocks.append({"shot": n, "cell": cells["时长(秒)"].coordinate, "field": "时长(秒)",
                           "code": "duration_invalid"})
        status = cells["来源/核对"]
        lines = str(status.value or "").split("\n")
        if lines[0] == "校验失败":
            blocks.append({"shot": n, "cell": status.coordinate, "field": "来源/核对", "code": "status_failed"})
        elif lines[0] == "待确认":
            warns.append({"shot": n, "cell": status.coordinate, "field": "来源/核对", "code": "status_pending"})
        for line in lines[1:]:
            if line and not line.startswith(("原文：", "AI建议：")):
                warns.append({"shot": n, "cell": status.coordinate, "field": "来源/核对",
                              "code": "status_note", "text": line})
    if n_shots == 0:
        blocks.append({"code": "no_shots"})
    return {"shots": n_shots, "blocked": bool(blocks), "blocks": blocks, "warnings": warns}


def main(out, files):
    out = Path(out)
    res = json.loads(out.read_text(encoding="utf-8")) if out.exists() else {}
    res.setdefault(GATE_KEY, {})
    for f in files:
        p = Path(f)
        shots = shots_of(p) or None
        res[p.name] = dryrun.run(p, shots)
        res[GATE_KEY][p.name] = gate(p)
    out.write_text(json.dumps(res, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2:])
