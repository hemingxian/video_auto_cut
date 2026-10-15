# -*- coding: utf-8 -*-
"""VideoAutoCut 标准分镜脚本模板（VAC-SB-1）生成器。

表头文字、顺序是给程序读的契约——改之前先看 模板设计说明.md 的「表头为什么这么定」。
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_VERSION = "VAC-SB-1"
SHEET_GUIDE = "填写说明"
SHEET_PROJECT = "项目信息"
SHEET_STORY = "分镜脚本"
SHEET_DICT = "字典"
MAX_ROWS = 200  # 下拉 / 文本格式预设到第几行

FONT_NAME = "微软雅黑"

# ---------------------------------------------------------------- 列定义
# (列名, 列宽, 必填, 软件现在读, 下拉字典键, 文本格式, 批注)
COLUMNS = [
    ("镜号", 6, "必填", True, None, False,
     "从 1 开始连续编号，一镜一行。不要跳号、不要写 1a / 片头。"),
    ("段落", 12, "选填", False, None, True,
     "章节名（2–6 字），同一段落的几镜写同一个名字。以后用作底部进度条的章节。"),
    ("分镜标题", 14, "选填", True, None, True,
     "这一镜的短标题。留空时软件自动叫「分镜N」。会逐字校验。"),
    ("时长(秒)", 8, "必填", True, None, False,
     "这一镜的目标时长，单位秒，写数字（如 8 或 7.5）。有口播时以实际配音长度为准。"),
    ("景别", 8, "选填", False, "景别", True, "下拉选择。见「填写说明」的景别定义。"),
    ("运镜", 8, "选填", False, "运镜", True, "下拉选择。见「填写说明」的运镜定义。"),
    ("拍摄角度", 9, "选填", False, "拍摄角度", True, "下拉选择。见「填写说明」的角度定义。"),
    ("画面内容", 30, "必填", True, None, True,
     "观众在这一镜看到什么：主体 + 动作 + 场景。写看得见的东西，不写观点。"),
    ("口播", 40, "有配音时必填", True, None, True,
     "旁白原文，逐字配音。没有配音的镜留空。"),
    ("字幕", 28, "选填", True, None, True,
     "屏幕字幕。留空 = 与口播相同。只有字幕要和口播不一样时才写。"),
    ("上屏文字", 30, "选填", True, None, True,
     "卡片上观众看到的字。没写 / 写「-」「无」= 这一镜不上卡片。\n"
     "行首标签（不上屏，只选版式）：大字： 副行： 小字： 参数卡： 决策树："),
    ("卡片提示", 14, "选填", True, "卡片提示", True,
     "卡片样式提示（给制作看，不上屏）。可下拉选，也可自己写。"),
    ("推荐素材来源", 24, "选填", True, None, True,
     "去哪找素材：素材库文件夹 / 指定文件（可带取段）/ 图库 / 客户提供。"),
    ("素材要求", 22, "选填", True, None, True,
     "素材的限制：横竖屏、清晰度、需打码的内容、不能出现的东西等。"),
    ("搜索关键词", 20, "选填", True, None, True,
     "给素材检索用的关键词，逗号分隔，中英文都可。留空则按画面内容检索。"),
    ("转场", 8, "选填", False, "转场", True, "进入下一镜的方式。留空 = 硬切。"),
    ("备注", 20, "选填", False, None, True, "音效、配乐、制作提醒等其它说明。"),
    ("来源/核对", 26, "选填", False, None, True,
     "作者直接手写的留空。由客户脚本转换来的，第一行写 已确认 / 待确认 / 校验失败，"
     "后面几行写原文出处与 AI 建议了哪些格子。"),
]
HEADERS = [c[0] for c in COLUMNS]
COL = {name: i + 1 for i, name in enumerate(HEADERS)}

# ---------------------------------------------------------------- 字典
DICT = {
    "景别": ["大远景", "远景", "全景", "中景", "中近景", "近景", "特写", "大特写", "不限"],
    "运镜": ["固定", "推", "拉", "左摇", "右摇", "上摇", "下摇", "左移", "右移", "跟",
             "升", "降", "环绕", "甩", "手持", "不限"],
    "拍摄角度": ["平视", "俯拍", "仰拍", "顶拍", "侧拍", "背拍", "主观视角", "倾斜", "不限"],
    "转场": ["硬切", "叠化", "淡入", "淡出", "闪白", "闪黑", "划像", "甩切", "匹配剪辑"],
    "卡片提示": ["标题卡", "编号卡", "要点列表", "参数列表", "流程条", "对照卡", "决策树",
               "行动号召"],
    "视频类型": ["产品宣传", "科普讲解", "案例展示", "企业宣传", "培训教学", "活动快剪", "其它"],
    "画幅": ["16:9", "9:16", "1:1"],
    "语言": ["中文", "英文", "中英双语"],
    "转换方式": ["作者手写", "列映射", "大模型抽取"],
}
DICT_HEADER = {  # 字典页表头（刻意不用任何分镜表会认的列名）
    "卡片提示": "卡片提示建议",
}

# ---------------------------------------------------------------- 样式
C_HEAD_READ = "1F3A5F"   # 软件现在读的列
C_HEAD_LATER = "5B6770"  # 先写着、以后用的列
C_SECTION = "E8EDF3"
C_AI = "FFF2CC"          # AI 建议，待人确认
C_BAD = "FFC7CE"         # 出处校验失败
C_BAND = "F7F9FB"
THIN = Side(style="thin", color="B7C0CA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_BODY = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_HEAD = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="1F3A5F")
F_H2 = Font(name=FONT_NAME, size=11, bold=True, color="1F3A5F")
F_NOTE = Font(name=FONT_NAME, size=9, color="5B6770")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


# ---------------------------------------------------------------- 字典页
def build_dict_sheet(wb: Workbook) -> dict[str, str]:
    ws = wb.create_sheet(SHEET_DICT)
    refs = {}
    for ci, (key, values) in enumerate(DICT.items(), 1):
        head = ws.cell(row=1, column=ci, value=DICT_HEADER.get(key, key))
        head.font, head.fill, head.alignment = F_HEAD, fill(C_HEAD_LATER), CENTER
        for ri, v in enumerate(values, 2):
            ws.cell(row=ri, column=ci, value=v).font = F_BODY
        letter = get_column_letter(ci)
        refs[key] = f"'{SHEET_DICT}'!${letter}$2:${letter}${len(values) + 1}"
        ws.column_dimensions[letter].width = 13
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "B7C0CA"
    return refs


# ---------------------------------------------------------------- 分镜表
def build_story_sheet(wb: Workbook, refs: dict[str, str], shots=None, marks=None,
                      title: str = SHEET_STORY):
    """shots: list[dict 列名->值]；marks: {(行序号0起, 列名): 'ai'|'bad'}。"""
    ws = wb.create_sheet(title)
    for ci, (name, width, req, reads, _dk, _tf, note) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=ci, value=name)
        c.font, c.alignment, c.border = F_HEAD, CENTER, BORDER
        c.fill = fill(C_HEAD_READ if reads else C_HEAD_LATER)
        c.comment = Comment(f"【{req}｜{'软件现在读' if reads else '先写着，以后用'}】\n{note}",
                            "模板")
        c.comment.width, c.comment.height = 260, 130
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30

    shots = shots or []
    marks = marks or {}
    last = MAX_ROWS
    for ri in range(2, last + 1):
        data = shots[ri - 2] if ri - 2 < len(shots) else None
        for ci, (name, _w, _r, _rd, _dk, textfmt, _n) in enumerate(COLUMNS, 1):
            c = ws.cell(row=ri, column=ci)
            if data is not None:
                v = data.get(name)
                if v not in (None, ""):
                    c.value = v
                c.border = BORDER
                c.font = F_BODY
                c.alignment = CENTER if name in ("镜号", "时长(秒)", "景别", "运镜", "拍摄角度", "转场") else WRAP_TOP
                m = marks.get((ri - 2, name))
                if m == "ai":
                    c.fill = fill(C_AI)
                elif m == "bad":
                    c.fill = fill(C_BAD)
                elif (ri % 2) == 1:
                    c.fill = fill(C_BAND)
            else:
                c.alignment = WRAP_TOP
                c.font = F_BODY
            if textfmt:
                c.number_format = "@"

    # 数据验证（只挂到空白模板行也生效；不会产生单元格值）
    def add_list(col_name: str, key: str, strict: bool):
        letter = get_column_letter(COL[col_name])
        dv = DataValidation(type="list", formula1=refs[key], allow_blank=True,
                            showErrorMessage=True,
                            errorStyle="stop" if strict else "information")
        dv.errorTitle = "请从下拉里选"
        dv.error = ("这一列只收字典里的值（见「填写说明」）。" if strict
                    else "不在建议值里——可以保留你自己的写法。")
        dv.promptTitle = col_name
        dv.prompt = "从下拉选择" + ("" if strict else "，也可以自己写")
        dv.showInputMessage = True
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)

    for col_name in ("景别", "运镜", "拍摄角度", "转场"):
        add_list(col_name, col_name, True)
    add_list("卡片提示", "卡片提示", False)

    letter = get_column_letter(COL["镜号"])
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="999",
                        allow_blank=True, showErrorMessage=True)
    dv.error, dv.errorTitle = "镜号写 1、2、3… 的整数", "镜号"
    dv.add(f"{letter}2:{letter}{last}")
    ws.add_data_validation(dv)
    letter = get_column_letter(COL["时长(秒)"])
    dv = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="600",
                        allow_blank=True, showErrorMessage=True)
    dv.error, dv.errorTitle = "时长写秒数，0.5–600 之间的数字", "时长(秒)"
    dv.add(f"{letter}2:{letter}{last}")
    ws.add_data_validation(dv)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = None
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.print_options.gridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "&A  第 &P / &N 页"
    ws.sheet_properties.tabColor = C_HEAD_READ
    return ws


# ---------------------------------------------------------------- 项目信息
PROJECT_ROWS = [
    # (键, 说明, 下拉)
    ("模板版本", "固定值，程序靠它认出这是标准模板，请勿修改", None),
    ("§基本信息", None, None),
    ("片名", "成片名称（不是分镜标题）", None),
    ("客户/品牌", "出片主体；无品牌露出的写「无品牌」", None),
    ("视频类型", "下拉选择", "视频类型"),
    ("目标平台", "如 抖音 / 视频号 / B站 / YouTube，可写多个", None),
    ("画幅", "下拉：16:9 / 9:16 / 1:1", "画幅"),
    ("目标时长(秒)", "整片目标时长；各镜时长之和见下一行", None),
    ("分镜合计时长(秒)", "自动计算（分镜脚本「时长(秒)」列之和），不用填", None),
    ("语言", "下拉选择", "语言"),
    ("§声音", None, None),
    ("配音音色", "如 男声·沉稳 / 女声·亲和，或具体音色名", None),
    ("配音语速", "如 正常 / 偏快，或 4.2 字/秒", None),
    ("BGM 风格", "如 轻电子、低音量、不抢口播；不要 BGM 写「无」", None),
    ("§包装（以后给品牌包用）", None, None),
    ("品牌主色", "十六进制色值，如 #17191C", None),
    ("品牌强调色", "十六进制色值，如 #FF7A33", None),
    ("字体", "如 思源黑体；没有要求留空", None),
    ("字幕样式", "如 描边 + 关键词高亮，不加底色条", None),
    ("角标/Logo", "是否常驻角标、放哪；不要写「无」", None),
    ("§素材与红线", None, None),
    ("默认素材范围", "默认到哪个素材库文件夹里找，逐镜可在「推荐素材来源」里另写", None),
    ("禁止出现", "全片不能出现的品牌、人物、文字、画面", None),
    ("§转换信息（由转换器填写；作者手写的模板只填「转换方式=作者手写」）", None, None),
    ("原始脚本文件名", "客户交来的原始脚本文件名", None),
    ("原始脚本 SHA-256", "原始文件的 SHA-256，用来证明核对的是同一份原文", None),
    ("转换时间", "YYYY-MM-DD HH:MM（北京时间）", None),
    ("转换方式", "作者手写 / 列映射 / 大模型抽取", "转换方式"),
    ("映射档案", "列映射时使用的客户映射档案名（可复用）", None),
    ("转换器版本", "", None),
    ("核对人", "在 Excel 里逐镜核对的人", None),
    ("核对完成时间", "全部镜改成「已确认」后填写", None),
]


def build_project_sheet(wb: Workbook, refs: dict[str, str], values: dict | None = None,
                        ai_keys: set | None = None):
    ws = wb.create_sheet(SHEET_PROJECT)
    values = dict(values or {})
    values.setdefault("模板版本", TEMPLATE_VERSION)
    values.setdefault("分镜合计时长(秒)", f"=SUM('{SHEET_STORY}'!D2:D{MAX_ROWS})")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 46
    for ci, h in enumerate(("项目", "填写", "说明"), 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font, c.fill, c.alignment, c.border = F_HEAD, fill(C_HEAD_LATER), CENTER, BORDER
    r = 2
    for key, note, dk in PROJECT_ROWS:
        if key.startswith("§"):
            c = ws.cell(row=r, column=1, value=key[1:])
            c.font, c.fill = F_H2, fill(C_SECTION)
            for ci in (2, 3):
                ws.cell(row=r, column=ci).fill = fill(C_SECTION)
            r += 1
            continue
        a = ws.cell(row=r, column=1, value=key)
        b = ws.cell(row=r, column=2, value=values.get(key))
        c = ws.cell(row=r, column=3, value=note)
        a.font, b.font, c.font = F_BOLD, F_BODY, F_NOTE
        for cell in (a, b, c):
            cell.border, cell.alignment = BORDER, WRAP_TOP
        if key not in ("分镜合计时长(秒)",):
            b.number_format = "@"
        if key in ("模板版本", "分镜合计时长(秒)"):
            b.fill = fill(C_SECTION)
        if ai_keys and key in ai_keys:
            b.fill = fill(C_AI)
        if dk:
            dv = DataValidation(type="list", formula1=refs[dk], allow_blank=True,
                                showErrorMessage=True, errorStyle="information")
            dv.add(f"B{r}")
            ws.add_data_validation(dv)
        r += 1
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "5B6770"
    return ws


# ---------------------------------------------------------------- 填写说明
TERMS = {
    "景别": [
        ("大远景", "EWS", "人或设备在画面里很小，交代大环境", "航拍整座厂区"),
        ("远景", "WS", "看得清场地和整条线，看不清细节", "整条产线从头到尾"),
        ("全景", "FS", "主体完整入画、四周留一点环境", "一台机床连同旁边的机器人完整入画"),
        ("中景", "MS", "主体的大半部分，动作看得清", "机器人从腰部到夹爪"),
        ("中近景", "MCU", "比中景更近，环境只剩一点", "夹爪和工件占画面一半"),
        ("近景", "CU", "局部占满画面，能看清结构", "夹爪抓着工件"),
        ("特写", "CU+", "一个细节占满画面", "卡盘夹紧工件那一下"),
        ("大特写", "ECU", "比特写更近，看纹理、字符", "工件表面的刀纹"),
        ("不限", "—", "交给软件按素材决定", "—"),
    ],
    "运镜": [
        ("固定", "Static", "机位不动", "三脚架定机位拍机器人循环"),
        ("推", "Push in", "镜头向主体靠近，画面逐渐放大", "从机床全景推到夹爪"),
        ("拉", "Pull out", "镜头远离主体，画面逐渐交代环境", "从工件特写拉到整台设备"),
        ("左摇/右摇", "Pan", "机位不动，镜头水平转动", "从机床摇到料仓"),
        ("上摇/下摇", "Tilt", "机位不动，镜头上下转动", "从地面料盘摇到机器人"),
        ("左移/右移", "Truck", "机位整体横向平移", "沿产线平移经过多台机床"),
        ("跟", "Follow", "镜头跟着运动主体走", "跟着 AGV 行进"),
        ("升/降", "Pedestal/Crane", "机位整体升高或降低", "从设备正面升到俯瞰"),
        ("环绕", "Orbit/Arc", "镜头绕主体转", "绕机器人工作站半圈"),
        ("甩", "Whip pan", "快速摇动带出模糊，常接转场", "甩到下一台设备"),
        ("手持", "Handheld", "轻微晃动的现场感", "车间走动视角"),
        ("不限", "—", "交给软件按素材决定", "—"),
    ],
    "拍摄角度": [
        ("平视", "Eye level", "镜头与主体同高", "正面平视机床操作面"),
        ("俯拍", "High angle", "从上往下斜着拍", "斜上方看双料盘"),
        ("仰拍", "Low angle", "从下往上拍，显得高大", "低机位仰拍重载机器人"),
        ("顶拍", "Top down", "正上方垂直向下", "正上方看料盘网格"),
        ("侧拍", "Side", "从主体侧面拍", "侧面看机械臂伸进卡盘"),
        ("背拍", "Back", "从主体后方拍", "机器人身后看机床"),
        ("主观视角", "POV", "模拟人或设备的视线", "从夹爪视角看工件靠近"),
        ("倾斜", "Dutch", "画面水平线倾斜，制造张力", "倾斜构图的故障画面"),
        ("不限", "—", "交给软件按素材决定", "—"),
    ],
    "转场": [
        ("硬切", "Cut", "直接切到下一镜（留空即硬切）", "—"),
        ("叠化", "Dissolve", "上一镜渐隐同时下一镜渐显", "时间流逝、同一场景前后"),
        ("淡入/淡出", "Fade", "从黑渐显 / 渐隐到黑", "片头、片尾"),
        ("闪白/闪黑", "Flash", "一帧白或黑的闪切", "节奏点、段落切换"),
        ("划像", "Wipe", "新画面从一侧划入", "对比前后两种方案"),
        ("甩切", "Whip", "配合甩镜的快速切换", "快节奏罗列"),
        ("匹配剪辑", "Match cut", "前后两镜形状/动作相接", "上一镜抓起、下一镜放下"),
    ],
}

LABEL_ROWS = [
    ("大字：", "这一行用大字号", "大字：抓起来、放进去？"),
    ("副行：/ 小字：", "这一行用小一号字，跟在大字后面", "副行：决定成败的其实是三件事"),
    ("参数卡：", "参数列表（写在第一行行首），暂按各行一样大出", "参数卡：总功率 <3 kW"),
    ("决策树：", "判断分支（写在第一行行首），暂按各行一样大出", "决策树：节拍 <60s → 料盘式单元"),
]
LAYOUT_RULES = [
    ("大字 + 副行/小字", "第一行大字，后面几行小一号", "第一行可以不写「大字：」"),
    ("大字 + 大字", "各行一样大", "每行都写「大字：」"),
    ("多行、都不写标签", "各行一样大（软件默认）", "—"),
    ("参数卡 / 决策树", "暂按各行一样大（以后单做版式）", "只在第一行写标签即可"),
    ("单行", "默认标题卡", "单行写不写标签都一样"),
]
FORBIDDEN = [
    "不要改表头文字、不要调换或删除列、不要在右侧新增列（识别会降档，见设计说明）。",
    "表头必须在第 1 行；表头和第 1 镜之间不要插说明行、片头行。",
    "不要合并单元格、不要隐藏行或列；一镜只占一行。",
    "镜号从 1 连续编到 N，不跳号、不重复，不写 1a、片头、小计。",
    "上屏文字、卡片提示、分镜标题三列不要用公式，也不要只写一个纯数字（需要时写成「60 秒」这类带字的文本）。",
    "同一个工作簿里只放一张分镜表：不要把这张表复制成第二个 sheet（中英两版请存两个文件）。",
    "行首标签后面必须有冒号（全角半角都行）：写成「大字 抓起来」会被当成拿不准，只提示不改。",
    "上屏文字里不要写 markdown 星号（**加粗**），画面上会原样显示星号。",
]
PLACEHOLDERS = "-  —  ——  /  ／  无  没有  暂无  空  留空  空白  不填  不上屏  N/A  （不填）  （无）"


def build_guide_sheet(wb: Workbook):
    ws = wb.active
    ws.title = SHEET_GUIDE
    widths = {"A": 16, "B": 12, "C": 12, "D": 54, "E": 40}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    r = 1

    def line(text, font=F_BODY, span=5, height=None, fill_color=None):
        nonlocal r
        c = ws.cell(row=r, column=1, value=text)
        c.font, c.alignment = font, WRAP_TOP
        if span > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
        if fill_color:
            c.fill = fill(fill_color)
        if height:
            ws.row_dimensions[r].height = height
        r += 1

    def table(head, rows, fills=None):
        nonlocal r
        for ci, h in enumerate(head, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.font, c.fill, c.alignment, c.border = F_HEAD, fill(C_HEAD_LATER), CENTER, BORDER
        r += 1
        for i, row in enumerate(rows):
            for ci, v in enumerate(row, 1):
                c = ws.cell(row=r, column=ci, value=v)
                c.font, c.alignment, c.border = (F_BOLD if ci == 1 else F_BODY), WRAP_TOP, BORDER
                if fills and fills.get(i):
                    c.fill = fill(fills[i])
            r += 1
        r += 1

    line(f"VideoAutoCut 标准分镜脚本模板（{TEMPLATE_VERSION}）· 填写说明", F_TITLE, height=26)
    line("按这份模板写的脚本，软件会直接按列读取、不经过 AI 改写。所以列名、写法都是固定的——"
         "照下面的规则写，你写什么，成片就用什么。", height=32)
    line("三条铁律：① 一切以脚本为准：口播、字幕、上屏文字、分镜标题逐字使用，软件不改写；"
         "② 上屏文字只在你写了的镜出现，没写就不上卡片；③ 不要改表头、不合并单元格、不隐藏行列。",
         F_BOLD, height=32, fill_color=C_SECTION)
    r += 1

    line("一、每一列怎么写", F_H2)
    line("表头深蓝色 = 软件现在就会读；表头灰色 = 先写着、以后的版本会用（现在写了也不会出错）。"
         "鼠标停在表头上可以看到简要提示。", F_NOTE, height=28)
    how = {
        "镜号": "1、2、3… 连续整数，一镜一行",
        "段落": "章节名，同段几镜写同一个名字；以后做底部进度条章节",
        "分镜标题": "这一镜的短标题；留空软件自动编「分镜N」",
        "时长(秒)": "数字，单位秒，可带一位小数；有口播时以实际配音为准",
        "景别": "下拉选；定义见下方第三节",
        "运镜": "下拉选；定义见下方第三节",
        "拍摄角度": "下拉选；定义见下方第三节",
        "画面内容": "看得见的东西：主体 + 动作 + 场景，一两句话",
        "口播": "旁白原文，逐字配音；这一镜不配音就留空",
        "字幕": "留空 = 与口播相同；只有要和口播不一样时才写",
        "上屏文字": "卡片上观众看到的字；可多行；行首标签选版式（见第二节）",
        "卡片提示": "卡片样式提示，不上屏；下拉选或自己写",
        "推荐素材来源": "素材库文件夹 / 指定文件（可写取段）/ 图库 / 客户提供",
        "素材要求": "横竖屏、清晰度、要打码或不能出现的内容",
        "搜索关键词": "逗号分隔，中英文都行；留空则按画面内容检索",
        "转场": "下拉选；留空 = 硬切",
        "备注": "音效、配乐、其它提醒",
        "来源/核对": "手写的留空；转换来的第一行写 已确认/待确认/校验失败（见第五节）",
    }
    example = {
        "镜号": "1", "段落": "钩子", "分镜标题": "开场钩子", "时长(秒)": "11",
        "景别": "全景", "运镜": "推", "拍摄角度": "平视",
        "画面内容": "机器人从机床里取出工件，镜头推近到夹爪",
        "口播": "机器人给机床上下料，看起来就是抓起来、放进去。",
        "字幕": "（多数情况留空）",
        "上屏文字": "大字：抓起来、放进去？\n副行：决定成败的其实是三件事",
        "卡片提示": "标题卡",
        "推荐素材来源": "素材库：机床上下料 / 指定：xx.mp4 取段 t60–t68",
        "素材要求": "横屏；铭牌和控制柜文字打码",
        "搜索关键词": "机器人 机床 取件, CNC tending",
        "转场": "叠化", "备注": "音效：气动夹爪咔嗒声",
        "来源/核对": "（作者手写时留空）",
    }
    table(("列名", "必填", "软件现在读", "怎么写", "例"),
          [(name, req, "读取" if reads else "暂不读取", how[name], example[name])
           for name, _w, req, reads, _dk, _tf, _n in COLUMNS])

    line("二、上屏文字的写法", F_H2)
    line("每一行开头可以写一个版式标签（带冒号）。标签本身不上屏，只用来选版式；"
         "不在行首的冒号（如「下一期：」「对接：」）照常上屏。", F_NOTE, height=28)
    table(("行首标签", "作用", "", "例", ""),
          [(a, b, "", c, "") for a, b, c in LABEL_ROWS])
    table(("标签组合", "出来的版式", "", "说明", ""),
          [(a, b, "", c, "") for a, b, c in LAYOUT_RULES])
    line("不上卡片的镜：上屏文字留空，或只写下面任一占位（整格只有它）："
         + PLACEHOLDERS + "。"
         "注意「无人化产线」这种正文不会被当成占位；「没有！」「——（不填）」这种半像占位的写法软件拿不准，"
         "会只提示不处理，请用上面的标准写法。", height=58)
    line("多行：在单元格里按 Alt+Enter（Mac：Option+Enter / Control+Option+Enter）换行，一行一个意思。", height=20)
    r += 1

    line("三、景别 / 运镜 / 拍摄角度 / 转场", F_H2)
    line("这几列现在软件还不读，但请尽量写：以后选素材会按景别、运镜去匹配素材标签，"
         "静态图片或太短的素材会按运镜自动做推拉摇移。拿不准就选「不限」或留空。", F_NOTE, height=30)
    for key, rows in TERMS.items():
        line(key, F_BOLD)
        table(("值", "英文", "", "定义", "一句话画面示意"),
              [(a, b, "", c, d) for a, b, c, d in rows])

    line("四、不要这样做（会让软件认不准，整份退回人工核对）", F_H2)
    for i, t in enumerate(FORBIDDEN, 1):
        line(f"{i}. {t}", height=20)
    r += 1

    line("五、「来源/核对」列与颜色（客户脚本由转换器转成本模板时使用）", F_H2)
    table(("标记", "含义", "", "导入时怎么处理", ""),
          [("浅黄底", "AI 建议的格子（景别、运镜、画面内容、关键词等制作字段）", "",
            "人看过、同意或改好后，把这一镜「来源/核对」第一行改成「已确认」", ""),
           ("浅红底", "逐字出处校验没通过：口播/字幕/上屏文字/标题在客户原文里找不到", "",
            "必须改成与原文一致或由客户确认；该镜不得进入正式导入", ""),
           ("无底色", "原文逐字搬运，或作者直接手写", "", "照常导入", "")],
          fills={0: C_AI, 1: C_BAD})
    line("「来源/核对」第一行只写三个词之一：已确认 / 待确认 / 校验失败；后面几行写原文出处"
         "（如「原文：xxx.xlsx › 分镜脚本 › 第 3 行」）和「AI建议：景别、运镜」。作者直接手写的模板这一列留空。",
         height=44)
    r += 1

    line("六、一镜的完整写法（竖排示意，完整示例见《分镜脚本示例_上下料_v1.xlsx》）", F_H2)
    demo = [
        ("镜号", "1"), ("段落", "钩子"), ("分镜标题", "开场钩子"), ("时长(秒)", "11"),
        ("景别", "全景"), ("运镜", "推"), ("拍摄角度", "平视"),
        ("画面内容", "机器人在机床前取出工件，镜头从机床全景慢慢推近到夹爪"),
        ("口播", "机器人给机床上下料，看起来就是抓起来、放进去。真正决定它能不能稳定跑一整天的，其实是三件事。"),
        ("字幕", "（留空 = 与口播相同）"),
        ("上屏文字", "大字：抓起来、放进去？\n副行：决定成败的其实是三件事"),
        ("卡片提示", "标题卡"),
        ("转场", "淡入"),
    ]
    for k, v in demo:
        a = ws.cell(row=r, column=1, value=k)
        b = ws.cell(row=r, column=2, value=v)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        a.font, b.font = F_BOLD, F_BODY
        a.border = b.border = BORDER
        a.alignment = b.alignment = WRAP_TOP
        if "\n" in v:
            ws.row_dimensions[r].height = 30
        r += 1

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "1F3A5F"
    return ws


def build_workbook(path, shots=None, project=None, marks=None, project_ai=None,
                   extra_example_sheet=None):
    wb = Workbook()
    build_guide_sheet(wb)
    # 字典要先建好才能引用；建完再挪到最后
    refs = build_dict_sheet(wb)
    build_project_sheet(wb, refs, project, project_ai)
    build_story_sheet(wb, refs, shots, marks)
    if extra_example_sheet is not None:  # 仅用于干跑对照：模板里再放一张「示例」表
        build_story_sheet(wb, refs, extra_example_sheet, None, title="示例")
    order = [SHEET_GUIDE, SHEET_PROJECT, SHEET_STORY] + (["示例"] if extra_example_sheet else []) + [SHEET_DICT]
    wb._sheets = [wb[n] for n in order]
    wb.active = wb.sheetnames.index(SHEET_STORY)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SHEET_STORY
    wb.save(path)
    return path
