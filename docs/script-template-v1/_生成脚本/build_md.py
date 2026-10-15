# -*- coding: utf-8 -*-
"""从入库的 示例_食品码垛_v1.xlsx 逐字转写出 Markdown 载体（0.5.16 步骤 0）。

产出（都在 docs/script-template-v1/ 下）：
- 示例_食品码垛_v1.md      与 xlsx 同内容的标准 Markdown 示例
- 分镜脚本模板_v1.md       空白 Markdown 模板（项目信息 + 2 个空分镜）
- md负例/*.md              4 份由示例改出来的负例/提醒例

读取口径与 0.5.15 模板导入器一致：单元格值只去首尾空白、CRLF/CR 统一成 LF（authored_field_value）；
「项目信息」只取 A、B 都非空的行，跳过表头行（项目|填写）和公式格（分镜合计时长，导入器也不读）。
写法按 docs/script-input-0.5.16-2026-10-08.md §4.1：单行值 `- 键：值`；多行值冒号后换行、每行缩进 2 空格。
用法：python3 build_md.py
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
XLSX = ROOT / "示例_食品码垛_v1.xlsx"
HEADERS = ["镜号", "段落", "分镜标题", "时长(秒)", "景别", "运镜", "拍摄角度", "画面内容",
           "口播", "字幕", "上屏文字", "卡片提示", "推荐素材来源", "素材要求", "搜索关键词",
           "转场", "备注", "来源/核对"]
INFO_KEYS = ["模板版本", "片名", "客户/品牌", "视频类型", "目标平台", "画幅", "目标时长(秒)", "语言",
             "配音音色", "配音语速", "BGM 风格", "品牌主色", "品牌强调色", "字体", "字幕样式",
             "角标/Logo", "默认素材范围", "禁止出现", "原始脚本文件名", "原始脚本 SHA-256",
             "转换时间", "转换方式", "映射档案", "转换器版本", "核对人", "核对完成时间"]


def authored(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\r\n?", "\n", str(value)).strip()


def field_lines(key: str, value: str) -> list[str]:
    if "\n" not in value:
        return [f"- {key}：{value}"]
    return [f"- {key}："] + [f"  {line}" for line in value.split("\n")]


def read_xlsx() -> tuple[list[tuple[str, str]], list[dict[str, str]]]:
    wb = load_workbook(XLSX)  # 不用 data_only：公式格要识别出来跳过
    info = []
    for row in wb["项目信息"].iter_rows(min_row=2, values_only=True):
        key, value = authored(row[0]), authored(row[1] if len(row) > 1 else None)
        if key and value and not value.startswith("="):
            info.append((key, value))
    sheet = wb["分镜脚本"]
    assert [authored(c.value) for c in sheet[1]][:18] == HEADERS
    shots = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = {HEADERS[i]: authored(row[i]) for i in range(18)}
        if any(values.values()):
            shots.append(values)
    return info, shots


def render(title: str, info: list[tuple[str, str]], shots: list[dict[str, str]]) -> str:
    lines = [f"# {title}", "", "## 项目信息"]
    for key, value in info:
        lines += field_lines(key, value)
    for shot in shots:
        lines += ["", f"## 分镜 {shot['镜号']}"]
        for key in HEADERS[1:]:
            if shot.get(key):
                lines += field_lines(key, shot[key])
    return "\n".join(lines) + "\n"


def blank_template() -> str:
    lines = ["# 片名", "",
             "（这一段是说明，导入时忽略。每镜一个「## 分镜 N」小节，N 从 1 连续；"
             "多行内容在冒号后换行、每行前空两格；不需要的字段可以整行删掉；"
             "字段名和可选值见 字段规范_v1.md。）", "", "## 项目信息", "- 模板版本：VAC-SB-1"]
    lines += [f"- {key}：" for key in INFO_KEYS[1:18]]
    for number in (1, 2):
        lines += ["", f"## 分镜 {number}"] + [f"- {key}：" for key in HEADERS[1:17]]
    return "\n".join(lines) + "\n"


def mutate(text: str, old: str, new: str) -> str:
    assert text.count(old) == 1, old
    return text.replace(old, new)


def main() -> None:
    info, shots = read_xlsx()
    title = dict(info)["片名"]
    example = render(title, info, shots)
    (ROOT / "示例_食品码垛_v1.md").write_text(example, encoding="utf-8")
    (ROOT / "分镜脚本模板_v1.md").write_text(blank_template(), encoding="utf-8")

    status = "- 来源/核对：\n  已确认\n  Mac 原创（2026-09-23），非客户稿转换；画面按 9/17 冻结食品码垛素材核对"
    neg = ROOT / "md负例"
    neg.mkdir(exist_ok=True)
    cases = {
        # 预期：阻断 shot_no_invalid（第 4 镜写成 5，之后全部错位）
        "负例1_镜号跳号_应阻断.md": mutate(example, "## 分镜 4\n", "## 分镜 5\n"),
        # 预期：阻断（同一镜同一字段写两次）
        "负例2_字段重复_应阻断.md": mutate(example, "- 运镜：推\n", "- 运镜：推\n- 运镜：拉\n"),
        # 预期：阻断 text_cell_pending（AI建议里含正文字段「口播」= 正文黄格）
        "负例3_正文待确认_应阻断.md": mutate(
            example, "- 备注：本镜不上卡片（上屏文字写「—」）；音效：输送带滚筒声\n" + status,
            "- 备注：本镜不上卡片（上屏文字写「—」）；音效：输送带滚筒声\n- 来源/核对：\n  待确认\n  AI建议：口播、景别"),
        # 预期：可导入；第 3 镜 ai_suggested=[shot_size, camera_angle] + status_pending 提醒
        "提醒例_制作字段待确认_可导入.md": mutate(
            example, "- 备注：音效：机械臂伺服声压在转身处\n" + status,
            "- 备注：音效：机械臂伺服声压在转身处\n- 来源/核对：\n  待确认\n  AI建议：景别、拍摄角度"),
    }
    for name, text in cases.items():
        (neg / name).write_text(text, encoding="utf-8")
    print(f"{len(shots)} shots, {len(info)} info keys -> {ROOT}")


if __name__ == "__main__":
    main()
