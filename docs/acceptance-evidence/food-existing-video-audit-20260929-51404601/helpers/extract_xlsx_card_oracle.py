from __future__ import annotations
import hashlib,json
from pathlib import Path
import openpyxl

AUDIT=Path(__file__).resolve().parent
XLSX=Path(r'D:\autocut\video-auto-cut\docs\script-template-v1\示例_食品码垛_v1.xlsx')
h=hashlib.sha256(XLSX.read_bytes()).hexdigest()
w=openpyxl.load_workbook(XLSX,read_only=True,data_only=True)
s=w['分镜脚本']
out=[]
for shot,row in ((1,2),(5,6),(8,9)):
    raw=s.cell(row,11).value
    hint=s.cell(row,12).value
    visible=[]
    if raw and raw!='—':
        for line in raw.splitlines():
            for prefix in ('大字：','副行：','参数卡：'):
                if line.startswith(prefix): line=line[len(prefix):]
            visible.append(line)
    out.append({'shot':shot,'xlsx_row':row,'raw_card_content':raw,'card_hint':hint,'expected_visible_lines':visible,'card_expected':bool(visible)})
receipt={'schema':'food-xlsx-card-oracle/v1','xlsx_path':str(XLSX),'xlsx_sha256':h,'sheet':'分镜脚本','card_content_column':'K 上屏文字','card_hint_column':'L 卡片提示','shots':out,'read_only':True}
(AUDIT/'receipts'/'xlsx-card-oracle.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'xlsx_sha256':h,'shots':out},ensure_ascii=False))
