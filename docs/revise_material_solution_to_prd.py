from __future__ import annotations

import io
import os
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from PIL import Image, ImageDraw, ImageFont


SOURCE = Path(r"D:\work\project\project\video-auto-cut-V2\docs\VideoAutoCut_三场景素材系统专业方案设计.docx")
OUTPUT = Path(
    r"D:\work\project\project\video-auto-cut-V2\docs"
    r"\VideoAutoCut_三场景素材系统专业方案设计_需求对齐修订版.docx"
)


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_paragraph(document: Document, old: str, new: str) -> None:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == old]
    if len(matches) != 1:
        raise RuntimeError(f"Paragraph match count for {old!r}: {len(matches)}")
    set_paragraph_text(matches[0], new)


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def set_row(document: Document, table_index: int, row_index: int, values: list[str]) -> None:
    row = document.tables[table_index].rows[row_index]
    if len(row.cells) != len(values):
        raise RuntimeError(
            f"T{table_index}R{row_index} expects {len(row.cells)} cells, got {len(values)}"
        )
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def replace_text_in_all_runs(document: Document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)


document = Document(SOURCE)

paragraph_replacements = {
    "文档用途：研发评审、接口评审、数据库设计与实施排期":
        "文档用途：研发评审、接口评审、数据库设计与实施排期｜需求基线：软件功能需求描述修订（2026-07-30）",
    "数据层不以文件路径作为素材身份。每个素材使用稳定 material_id，每个原始内容使用 content_id，通过 asset_version 和 source_sha256 识别版本与内容；本地状态由当前设备检测，云端状态由服务端返回，两个状态组合生成“仅本地、仅云端、本地和云端、下载中、同步失败”等页面标识。":
        "数据层不以文件路径作为素材身份。material_id 表示某个用户素材库中的一条素材；asset_id + asset_version 表示一份确定内容的原视频，source_sha256 用于上传校验和内容去重。origin 属于原视频资产，added_from 属于用户素材条目，二者不能混用。本地状态由当前设备检测，云端状态由服务端返回，组合生成“仅本地、仅云端、本地和云端、上传中、下载中、同步失败”等页面标识。",
    "•\t本地素材、平台下载素材和云端恢复素材统一进入托管体系，并保留来源和授权属性。":
        "•\t外部文件、平台下载素材和云端恢复素材统一进入托管素材库；用户导入默认仅本地，平台下载按服务端启用规则自动同步。",
    "•\t加载任务当前修订版本、分镜顺序、文案、画面描述、剪辑要求和素材绑定。":
        "•\t从当前设备的本地 SQLite 加载任务草稿、分镜顺序、文案、画面描述、剪辑要求和素材绑定。",
    "1.\t打开任务时读取本地 SQLite 草稿；存在云端草稿时比较 revision 并执行同步。":
        "1.\t打开任务时只读取当前设备的本地 SQLite 草稿；第一版不提供云端草稿或跨设备草稿恢复。",
    "2.\t用户修改字段后先更新本地事务，再将操作写入 local_sync_outbox。":
        "2.\t用户修改字段后更新本地草稿事务；素材、文件夹、上传和下载等云端操作才写入 local_sync_outbox，任务草稿不进入云同步队列。",
    "3.\t自动保存以 650 至 1000 毫秒防抖提交，服务端使用 expected_revision 做乐观锁。":
        "3.\t自动保存以 650 至 1000 毫秒防抖提交到本地 SQLite，手动保存和自动保存更新同一份草稿并递增本地 revision。",
    "4.\t服务端保存成功后返回新 revision，本地清除 dirty 标记。":
        "4.\t本地事务提交成功后更新“已保存”和最后保存时间；失败时保留页面内容并提供重试入口。",
    "5.\t发现版本冲突时保留本地草稿，提示重新加载或基于最新版本合并，不能静默覆盖。":
        "5.\t同一设备出现并发保存时使用本地 revision 检测冲突，保留较新草稿并明确提示，不能静默覆盖。",
    "6.\t点击开始渲染时，服务端或本地编排生成 render_material_snapshot。":
        "6.\t点击开始生成时，以当时的 task_revision 创建新的 run_id，展开文件夹并生成不可变的 render_material_snapshot。",
    "•\t视频绑定和文件夹绑定合并后按 material_id + asset_version 去重。":
        "•\t选择器确认阶段只保存 material_id 或 folder_id；开始生成并展开文件夹后，按 asset_id + asset_version 合并去重。",
    "编辑阶段保存 folder_id 和 include_subfolders，便于用户继续维护文件夹。点击开始渲染时展开文件夹，生成不可变的素材清单快照。之后新增或删除文件只影响下一次渲染，不影响正在运行的任务。":
        "编辑阶段只保存 folder_id 和 include_subfolders。点击开始生成时展开文件夹，解析 material_id 指向的 asset_id + asset_version，并把 material_id、asset_id、asset_version、source_sha256 写入不可变清单。之后新增、删除、移动或改名只影响下一次生成，不改变当前 run_id。",
    "•\t“我的素材”来自托管素材目录和本地 SQLite 的素材记录。":
        "•\t“我的素材”来自配置数据根目录下的托管素材目录和本地 SQLite；用户在界面中只看到素材名称，不依赖原始外部路径。",
    "三个页面不分别建库。数据库按照数据性质和负载特征拆分：事务业务数据进入关系库，高维向量进入向量库，当前设备状态进入本地 SQLite。原视频、封面和预览文件进入腾讯云 COS。":
        "三个页面不分别建库。数据库按照数据性质和负载特征拆分：事务业务数据进入关系库，高维向量进入向量库，当前设备状态与任务草稿进入本地 SQLite。需要云端备份的原视频进入腾讯云 COS；对象存储不承载业务元数据、草稿或向量。",
    "4.\t内容一致则复用 content_id 和已有分析结果；内容不同则增加 collision_no。":
        "4.\t内容一致时复用已有 asset_id + asset_version 及分析结果；内容不一致时创建新的原视频资产，不能覆盖已有对象。",
    "6.\tobject_key 使用随机 content_id 和版本，避免哈希碰撞覆盖对象。":
        "6.\tobject_key 由服务端按不可预测的 asset_id 和版本生成并保持内部可见，不能以客户端哈希直接命名，也不能返回桌面端持久化。",
    "向量数据库不保存用户文件夹树、任务草稿或本地路径。正常素材列表只查询业务关系库或本地读模型，仅在语义搜索和分镜选段时访问向量库。":
        "第一版固定使用当前 qwen3-vl-embedding、768 维向量和一套服务端管理的 analysis_profile_id；只有模型、维度、切片和预处理规则完全一致的分析结果才允许复用。向量数据库不保存用户文件夹树、任务草稿或本地路径，正常素材列表只查询业务关系库或本地读模型，仅在语义搜索和分镜选段时访问向量库。",
    "1.\t按 material_id + asset_version 查询 local_material_file_map。":
        "1.\t先由 material_id 解析当前 asset_id + asset_version，再查询当前设备的 local_material_file_map 和 local_asset_file。",
    "页面可以隐藏真实路径，但主进程必须知道每个文件的相对路径。手动导入素材和平台下载素材可以位于不同子目录，只要它们共享稳定 material_id、content_id 和本地映射即可。":
        "页面可以隐藏真实路径，但 Electron 主进程必须保存相对于配置数据根目录的托管路径。用户导入、平台下载和云端恢复可位于不同子目录；material_id 指向用户条目，asset_id + asset_version 指向实际原视频，本机映射只存在当前设备。",
    "•\t所有写接口携带 Idempotency-Key，服务端对相同用户和操作键返回同一结果。":
        "•\t所有云端写接口携带 Idempotency-Key，服务端对相同用户和操作键返回同一结果；本地草稿保存使用本地事务和 revision。",
    "•\t更新任务、素材和文件夹携带 expected_revision，冲突返回明确的 409。":
        "•\t更新云端素材和文件夹时携带 expected_revision，冲突返回明确的 409；第一版任务草稿不调用云端更新接口。",
    "7.3 任务、分镜与绑定接口":
        "7.3 本地草稿、分镜校验与生成清单接口",
    "8.2 本地导入事务":
        "8.2 托管导入与上传策略",
    "1.\t文件先复制到目标目录中的临时名称，同时计算哈希。":
        "1.\tWindows 选择的文件若不在托管工作目录，Electron 后台先复制到配置数据根目录中的临时文件，并在复制过程中计算 SHA-256 和大小；不能直接保存用户原始路径。",
    "2.\t复制完成并校验后原子重命名，避免半文件出现在素材库。":
        "2.\t复制完成并校验后原子重命名，再创建或复用 asset_id + asset_version、material_id 和本机路径映射；校验完成前不能绑定到分镜。",
    "3.\t同一 SQLite 事务写入素材、文件映射和分析任务。":
        "3.\t同一 SQLite 事务写入素材条目、原视频映射和 VL 分析任务；内容完全相同时复用原视频资产和已有分析，不重复保存物理文件。",
    "4.\t事务提交后通知页面刷新；失败时删除临时文件并保留错误原因。":
        "4.\t用户导入素材默认仅本地，“自动备份我导入的素材”默认关闭，也可手动上传；平台下载素材按服务端启用规则自动上传。上传是独立可恢复任务，失败不影响已完整导入的本地素材。",
    "•\t任务草稿、素材和文件夹全部使用 revision 乐观锁。":
        "•\t任务草稿使用当前设备的本地 revision；云端素材和文件夹使用 expected_revision 乐观锁，二者不做跨设备草稿合并。",
    "•\t临时对象放在独立前缀并配置生命周期清理，正式对象按 content_id 和版本组织。":
        "•\t临时对象放在独立前缀并配置生命周期清理，正式对象按服务端 asset_id 和版本组织；object_key 不返回客户端。",
    "1.\t先建立稳定素材ID、内容版本和本地SQLite，不在路径模型上直接叠加云同步。":
        "1.\t先建立稳定 shot_id、material_id、folder_id、asset_id + asset_version 和本地 SQLite，不在路径模型上直接叠加云同步。",
    "•\t使用功能开关分别控制本地清单、分镜绑定、用户上传和平台自动同步。":
        "•\t用户导入素材的自动备份开关默认关闭，关闭只影响后续导入且不删除已上传内容；平台下载自动同步由服务端允许名单控制。",
    "•\t先对测试账号开放云同步，观察上传成功率、哈希校验、重试和COS成本。":
        "•\t先对测试账号开放云备份，观察上传成功率、哈希校验、断点恢复、VL同步和 COS 成本；自动上传不等于自动公开。",
    "•\t数据结构迁移向前兼容；旧路径素材在后台逐步补 material_id 和 SHA-256。":
        "•\t旧路径素材在后台逐步迁入托管目录并补齐 material_id、folder_id、asset_id、asset_version 和 SHA-256，不在启动时一次性阻塞处理。",
    "•\t刷新和重启后可以恢复任务草稿、分镜顺序和素材绑定。":
        "•\t刷新和重启后可以恢复当前设备上的任务草稿、分镜顺序和素材绑定；换设备不承诺恢复草稿。",
    "•\t确认结果保存稳定ID、版本和哈希，不保存本机路径。":
        "•\t确认结果只保存 material_id 或 folder_id + include_subfolders，不保存版本、哈希、本机路径、object_key 或临时 URL。",
    "•\t任务草稿第一版是否只要求当前设备恢复。":
        "•\t上传会话有效期、用户容量、并行上传数、删除保留期和自动同步平台名单的部署取值。",
    "•\tmaterial_id 是业务身份，source_sha256 不是业务主键。":
        "•\tmaterial_id 是用户素材条目身份，asset_id + asset_version 是原视频身份，source_sha256 不是业务主键。",
    "•\t分镜绑定在渲染开始时固化不可变快照。":
        "•\t分镜草稿只保存 material_id / folder_id 绑定，开始生成时才固化包含 asset_id、asset_version 和 source_sha256 的不可变清单。",
    "•\t用户私有素材默认不做可感知的跨用户去重。":
        "•\t跨用户私有素材只允许服务端内部做物理去重，普通客户端不能查询哈希命中或感知其他用户是否上传过同一视频。",
}

for old_text, new_text in paragraph_replacements.items():
    replace_paragraph(document, old_text, new_text)

set_row(document, 0, 1, ["V1.1", "需求对齐修订稿", "2026-07-31", "桌面端、云素材服务、生成编排"])
set_row(document, 2, 3, ["素材身份", "material_id + asset_id + asset_version", "用户条目与原视频身份分离，路径、名称和临时URL不参与识别"])
set_row(document, 2, 4, ["内容去重", "SHA-256 + 文件大小候选；服务端二次校验", "常用索引保持紧凑，碰撞不覆盖原对象"])
set_row(document, 4, 1, ["素材身份", "主要依赖本地路径和文件信息", "稳定 material_id、folder_id、asset_id + asset_version"])
set_row(
    document,
    5,
    0,
    ["强约束：任务草稿只保存 material_id，或 folder_id + include_subfolders。草稿不保存 asset_version、哈希、本机路径、object_key、临时 URL、名称快照或排序；这些内容只在开始生成时写入本次运行清单。"],
)
set_row(document, 6, 1, ["模式一致性", "download_material=true 但存在绑定，或 false 且生成前无有效绑定", "切换为自动下载时确认清空绑定；使用已选素材时补齐有效绑定"])
set_row(document, 6, 5, ["绑定解析", "material_id 或 folder_id 已删除、无权访问或展开后候选为0", "返回 shot_id、material_id 和可执行处理动作"])
set_row(document, 7, 1, ["单个/多个视频", "仅 material_id", "同一分镜去重后写入素材绑定"])
set_row(document, 7, 4, ["仅云端素材", "仅 material_id", "生成前解析版本并下载校验，再允许参与生成"])
set_row(document, 9, 1, ["业务关系数据库", "1套", "服务端", "用户素材条目、文件夹、原视频资产、上传会话、分析、公开记录、生成运行清单", "是"])
set_row(document, 9, 5, ["腾讯云 COS", "1个私有桶/环境", "云端", "原视频和临时上传对象", "文件可信源"])
set_row(document, 10, 5, ["任务草稿", "当前设备本地 SQLite（第一版）", "不能上传到云端或只保存在React内存"])

set_row(document, 11, 1, ["material_id", "某个用户素材库中的一条素材", "重命名、移动后不变；同一原视频可被不同用户条目引用"])
set_row(document, 11, 2, ["asset_id", "原视频资产内部编号", "与路径、文件名、平台地址无关"])
set_row(document, 11, 3, ["asset_version", "原视频资产的确定版本", "与 asset_id 一起确定内容；第一版不同内容创建新资产"])
set_row(document, 11, 4, ["source_sha256", "上传完成校验与内容去重候选", "二进制32字节，不直接作为业务主键或对象名"])
set_row(document, 11, 5, ["folder_id", "逻辑文件夹内部编号", "创建时生成一次，改名、移动后不变"])
set_row(document, 11, 6, ["revision", "用户条目、文件夹或本地草稿的更新版本", "各自作用域内递增，不能把本地草稿 revision 当成云端草稿版本"])

core_rows = [
    ["material_folder", "folder_id", "owner_id、parent_folder_id、name、revision", "与名称和物理路径无关的用户文件夹树"],
    ["material", "material_id", "owner_id、folder_id、asset_id、asset_version、display_name、added_from", "用户素材条目；不保存本机路径"],
    ["asset", "asset_id", "origin、mime_type、原始元数据、创建时间", "不可变原视频资产；origin 只属于本表"],
    ["asset_version", "asset_id+asset_version", "source_sha256、大小、object_storage_status、object_key", "确定内容版本；object_key 仅服务端可见"],
    ["analysis", "analysis_id", "asset_id、asset_version、analysis_profile_id、status", "原视频资产拥有的 VL 分析"],
    ["material_publication", "publication_id", "material_id、权利确认、状态、下架信息", "公开状态与云端备份状态分离"],
    ["upload_session", "upload_session_id", "owner_id、asset_id、临时对象、进度、过期时间", "短期权限、断点恢复和完成校验"],
    ["generation_run", "run_id", "task_id、task_revision、status、created_at", "一次生成运行；重试沿用原清单"],
    ["render_material_snapshot", "run_id+shot_id", "download_material、created_at、checksum", "每个分镜不可变的本次生成素材范围"],
    ["render_material_snapshot_item", "run_id+shot_id+material_id", "asset_id、asset_version、source_sha256", "文件夹展开后的具体素材；按资产版本去重"],
    ["outbox_event", "event_id", "event_type、payload、retry、status", "跨业务库、COS和向量库的最终一致性"],
    ["audit_log", "audit_id", "actor、action、resource、result", "敏感操作和权限审计"],
]
for index, values in enumerate(core_rows, start=1):
    set_row(document, 12, index, values)

local_rows = [
    ["local_material", "material_id、cloud_revision、folder_id、asset_id、asset_version、状态", "统一素材库页面的当前设备读模型"],
    ["local_asset_file", "local_file_id、asset_id、asset_version、hash、size、relative_path、status", "托管原视频在当前设备的真实文件清单"],
    ["local_material_file_map", "material_id、asset_id、asset_version、local_file_id", "用户素材条目与当前设备文件映射"],
    ["local_task_draft", "task_id、schema_version、revision、snapshot_json、save_status", "第一版任务草稿、自动保存和重启恢复"],
    ["local_sync_outbox", "operation_id、素材/文件夹/上传类型、payload、retry", "只同步云素材相关操作，不同步任务草稿"],
    ["local_download_session", "material_id、asset_id、asset_version、temp_path、bytes、status", "断点下载、校验和原子落盘"],
]
for index, values in enumerate(local_rows, start=1):
    set_row(document, 13, index, values)

set_row(
    document,
    15,
    0,
    ["analysis_id / asset_id / asset_version / analysis_profile_id / segment_id / start_ms / end_ms / embedding_model / embedding_dimension / slice_rule_version / preprocessing_version / embedding"],
)
set_row(document, 17, 1, ["用户私有素材", "服务端内部资产范围", "允许物理复用，不对客户端暴露命中", "客户端不能用哈希探测其他用户私有文件"])
set_row(
    document,
    19,
    0,
    ["<dataRoot>/video/materials/imported/...\n<dataRoot>/video/materials/platform/...\n<dataRoot>/video/materials/cloud/<asset_id>/<asset_version>/...\n<dataRoot>/video/runs/<run_id>/raw/..."],
)
set_row(document, 20, 7, ["DELETE", "/v1/materials/{id}/cloud-copy", "删除云端备份", "服务端只处理云端副本；本地删除由 Electron 执行"])

set_row(document, 21, 1, ["IPC", "taskDraft.load / taskDraft.save", "读取、自动保存和手动保存本机草稿", "仅当前设备 SQLite，使用 schema_version 与 revision"])
set_row(document, 21, 2, ["POST", "/v1/shot-material-validations", "校验分镜绑定、权限、云端状态和VL状态", "一次返回全部阻断项和处理动作"])
set_row(document, 21, 3, ["POST", "/v1/generation-runs", "按 task_revision 创建 run_id 并固化清单", "展开文件夹并按 asset_id + asset_version 去重"])
set_row(document, 21, 4, ["GET", "/v1/generation-runs/{run_id}/material-manifest", "读取本次生成素材清单", "同一 run_id 重试必须返回同一清单"])
set_row(document, 21, 5, ["POST", "/v1/analyses/status-checks", "批量检查或触发所需 VL 分析", "analysis_profile_id 由服务端控制"])
set_row(document, 21, 6, ["GET", "/v1/generation-runs/{run_id}", "查询运行状态", "阶段、进度、错误包含 shot_id 和 material_id"])

set_row(
    document,
    22,
    0,
    ['{\n  "material_id": "mat_01...",\n  "display_name": "产品演示_主镜头.mp4",\n  "folder_id": "fld_01...",\n  "added_from": "user_import",\n  "asset": {\n    "asset_id": "asset_01...",\n    "asset_version": 1,\n    "source_sha256": "4f...",\n    "origin": "user_upload"\n  },\n  "cloud_status": "ready",\n  "analysis_status": "ready",\n  "publication_status": "private",\n  "revision": 12\n}'],
)
set_row(
    document,
    23,
    0,
    ["托管目录导入并创建/复用 asset 与 material\n  -> 根据来源和用户策略决定是否上传\n  -> 创建上传会话并获取短期权限\n  -> 上传 COS 临时对象\n  -> 服务端校验大小和 SHA-256\n  -> 原子切换 asset_version 的 object_key 与 cloud_status\n  -> Outbox 同步分析记录和向量状态"],
)
set_row(document, 24, 4, ["VL分析", "pending -> analyzing", "是", "ready / failed"])
set_row(document, 24, 5, ["向量同步", "pending -> syncing", "是", "ready / failed"])

phase_rows = [
    ["阶段一：分镜选材正确", "shot_id、download_material、两类绑定、本次生成素材清单、限定范围检索", "用户选什么，生成只使用什么", "Rust、Python/VL、Electron"],
    ["阶段二：本地素材库与草稿", "托管导入、稳定ID、SQLite、文件操作、VL、草稿和制作中预览", "重启可恢复本机任务与素材状态", "Electron主进程与页面"],
    ["阶段三：云备份与跨设备素材恢复", "COS、上传会话、断点恢复、短期授权、状态合并、云分析同步", "用户素材按开关上传，平台素材按服务端规则自动同步", "云素材服务与鉴权"],
    ["阶段四：公共素材", "权利确认、publication、搜索、添加、取消公开、暂停和下架", "公开与备份分离，复用资产和分析", "合规审核与运营能力"],
]
for index, values in enumerate(phase_rows, start=1):
    set_row(document, 26, index, values)

set_row(document, 27, 1, ["断网编辑任务", "本地草稿正常保存；恢复网络后不上传草稿，只继续素材相关队列"])
set_row(document, 27, 5, ["同名不同视频", "创建不同 asset_id，不按文件名或路径合并"])
set_row(document, 27, 6, ["相同视频不同文件名", "复用 asset_id + asset_version 和VL分析，保留各自用户素材名称"])
set_row(document, 27, 7, ["SHA-256候选冲突", "服务端执行二次校验，创建新资产且不覆盖已有COS对象"])

set_row(document, 28, 0, ["事项", "已确定或建议", "实施影响"])
set_row(document, 28, 1, ["任务草稿范围", "第一版仅当前设备保存，不做云同步", "素材跨设备恢复与草稿恢复严格分离"])
set_row(document, 28, 2, ["平台素材自动同步", "只对服务端启用的平台生效；自动上传不等于公开", "平台名单、授权和隐私说明必须可配置"])
set_row(document, 28, 3, ["用户私有素材跨用户去重", "仅服务端内部做物理去重，客户端不可探测命中", "权限查询先按用户隔离，哈希接口不得泄露存在性"])
set_row(document, 28, 4, ["云向量数据库选型", "按768维、过滤条件、数据量和QPS压测决定", "analysis_profile_id 由服务端固定管理"])
set_row(document, 28, 5, ["旧素材迁移", "后台渐进迁入托管目录并补稳定ID和哈希", "不能阻塞启动，失败可重试"])
set_row(document, 28, 6, ["公开素材审核与下架", "独立 publication 能力，云备份不依赖公开", "权利确认、举报、暂停与下架需要独立记录"])
set_row(
    document,
    29,
    0,
    ["最终建议：严格按需求阶段推进。先完成分镜级 download_material、稳定 shot_id、material_id、folder_id 和生成运行清单，再完善本地托管素材库与当前设备草稿；随后接入云端业务库、腾讯云 COS 和分析同步，最后建设公开素材。不能在旧路径身份或任务级下载开关上直接叠加云同步。"],
)

status_rows = [
    ["import_status", "copying / verifying / ready / failed"],
    ["local_status", "present / downloading / missing / damaged"],
    ["analysis_status", "pending / analyzing / ready / failed"],
    ["cloud_status", "not_uploaded / uploading / verifying / ready / failed / deleted"],
    ["vector_sync_status", "pending / syncing / ready / failed"],
    ["publication_status", "private / publishing / public / unpublishing / suspended / removed"],
]
for index, values in enumerate(status_rows, start=1):
    set_row(document, 30, index, values)

index_rows = [
    ["asset_version", "INDEX(source_sha256, source_size)", "内容候选；不是业务唯一键"],
    ["material", "(owner_id, folder_id, deleted_at, updated_at DESC)", "文件夹素材分页"],
    ["material", "(owner_id, cloud_status, updated_at DESC)", "存储状态筛选"],
    ["material_folder", "UNIQUE(owner_id, parent_folder_id, normalized_name)", "同级名称约束，folder_id 与路径无关"],
    ["analysis", "UNIQUE(asset_id, asset_version, analysis_profile_id)", "相同分析配置只保留一份有效结果"],
    ["render_material_snapshot_item", "UNIQUE(run_id, shot_id, asset_id, asset_version)", "本次生成清单按实际原视频去重"],
    ["local_task_draft", "UNIQUE(task_id)", "当前设备每个任务一份活动草稿"],
    ["shot_folder_binding", "UNIQUE(shot_id, folder_id)", "防止同一分镜重复绑定文件夹"],
    ["outbox_event", "(status, next_retry_at)", "后台任务拉取"],
    ["local_asset_file", "(source_sha256, source_size)", "当前设备重复内容候选"],
]
for index, values in enumerate(index_rows, start=1):
    set_row(document, 31, index, values)

# Remove any stale terminology that is no longer valid in the aligned requirements.
replace_text_in_all_runs(document, "content_id", "asset_id")

document.core_properties.title = "VideoAutoCut 云端素材库三场景方案设计（需求对齐修订版）"
document.core_properties.subject = "依据《软件功能需求描述修订-20260730》修订"
document.core_properties.comments = (
    "V1.1：统一 material_id 与 asset_id + asset_version 身份边界；"
    "第一版任务草稿仅当前设备；生成时固化素材清单；补充托管导入与上传策略。"
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)


def redraw_architecture_image(image_bytes: bytes) -> bytes:
    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    draw = ImageDraw.Draw(image)
    bold_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 28)
    body_font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 20)

    def centered_text(box, text, font, fill):
        left, top, right, bottom = box
        bounds = draw.textbbox((0, 0), text, font=font)
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        x = left + (right - left - width) / 2
        y = top + (bottom - top - height) / 2 - bounds[1]
        draw.text((x, y), text, font=font, fill=fill)

    # Electron local-domain service box.
    draw.rectangle((430, 337, 1170, 410), fill="#f2f4f8")
    centered_text((430, 337, 1170, 374), "Electron 主进程 / 本地领域服务", bold_font, "#17365d")
    centered_text((430, 374, 1170, 410), "文件操作、草稿、本机映射、上传下载队列", body_font, "#64748b")

    # Cloud material and generation API box.
    draw.rectangle((600, 527, 1000, 625), fill="#fff9e9")
    centered_text((600, 532, 1000, 576), "云素材与生成 API", bold_font, "#17365d")
    centered_text((600, 576, 1000, 618), "鉴权、业务校验、幂等与短期授权", body_font, "#64748b")

    # Relational database and COS descriptions.
    draw.rectangle((75, 775, 365, 845), fill="#f7f9fb")
    centered_text((75, 775, 365, 810), "业务关系数据库", bold_font, "#17365d")
    centered_text((75, 810, 365, 845), "素材、文件夹、资产、运行清单", body_font, "#64748b")

    draw.rectangle((455, 775, 745, 845), fill="#f7f9fb")
    centered_text((455, 775, 745, 810), "腾讯云 COS", bold_font, "#17365d")
    centered_text((455, 810, 745, 845), "原视频与临时上传对象", body_font, "#64748b")

    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


with ZipFile(OUTPUT, "r") as source_zip:
    architecture = redraw_architecture_image(source_zip.read("word/media/image1.png"))
    temp_path = OUTPUT.with_suffix(".tmp.docx")
    with ZipFile(temp_path, "w", ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            data = architecture if item.filename == "word/media/image1.png" else source_zip.read(item.filename)
            target_zip.writestr(item, data)
os.replace(temp_path, OUTPUT)
print(OUTPUT)
